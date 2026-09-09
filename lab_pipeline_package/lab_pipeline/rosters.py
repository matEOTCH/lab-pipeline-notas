"""Reading student lists and writing the clean Lista/Notas workbook.

Two inputs produce students: the raw Blackboard course-activity CSV (used once,
when a section is first created) and the clean Lista workbook this module writes
(used by every later workflow).
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .common import (
    clean_text,
    decode_csv_bytes,
    find_col_ws,
    find_header_row_ws,
    normalize_header,
    normalize_student_code,
)


LISTA_HEADERS = ["Nro.", "Codigo", "Nombre Completo", "Email"]
LISTA_COLUMN_WIDTHS = {"A": 8, "B": 14, "C": 48, "D": 34}

HEADER_FILL = "FF7F00"
BORDER_COLOR = "B7B7B7"

REQUIRED_ROSTER_COLUMNS = {
    "apellido": "Apellido",
    "nombre": "Nombre",
    "nombre de usuario": "Nombre de usuario",
}


def _sniff_dialect(sample: str):
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t")
    except Exception:
        return csv.excel


def read_delimited(path: Path) -> list[dict[str, str]]:
    """Read a CSV whose delimiter we do not know in advance."""
    text = decode_csv_bytes(path)
    handle = io.StringIO(text, newline="")
    return list(csv.DictReader(handle, dialect=_sniff_dialect(text[:4096])))


# --- Blackboard course-activity export ----------------------------------------


def read_blackboard_roster(csv_path: Path) -> list[dict[str, str]]:
    """Blackboard's course-activity CSV -> sorted students.

    Rows missing any of surname / given name / username are skipped: the export
    carries footer and summary rows that are not students.
    """
    rows = read_delimited(csv_path)
    if not rows:
        raise ValueError("El CSV esta vacio.")

    columns = {normalize_header(column): column for column in rows[0].keys()}
    missing = [label for key, label in REQUIRED_ROSTER_COLUMNS.items() if key not in columns]
    if missing:
        raise ValueError(f"Faltan columnas requeridas en el CSV: {', '.join(missing)}")

    students = []
    for row in rows:
        apellido = clean_text(row.get(columns["apellido"])).upper()
        nombre = clean_text(row.get(columns["nombre"])).upper()
        codigo = normalize_student_code(row.get(columns["nombre de usuario"]))
        if not apellido or not nombre or not codigo:
            continue
        students.append({"apellido": apellido, "nombre": nombre, "codigo": codigo})

    students.sort(key=lambda item: (item["apellido"], item["nombre"], item["codigo"]))
    return students


def build_clean_list(students: list[dict[str, str]], email_domain: str) -> list[dict[str, str]]:
    domain = clean_text(email_domain).removeprefix("@")
    return [
        {
            "Nro.": index,
            "Codigo": student["codigo"],
            "Nombre Completo": f"{student['apellido']}, {student['nombre']}",
            "Email": f"{student['codigo']}@{domain}",
        }
        for index, student in enumerate(students, start=1)
    ]


def save_students_workbook(rows: list[dict[str, str]], out_path: Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Lista"

    ws.append(LISTA_HEADERS)
    for row in rows:
        ws.append([row[header] for header in LISTA_HEADERS])

    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for column, width in LISTA_COLUMN_WIDTHS.items():
        ws.column_dimensions[column].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
    for row in range(2, ws.max_row + 1):
        for column in (1, 2):
            ws.cell(row=row, column=column).alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(out_path)
    return out_path


# --- the clean Lista workbook -------------------------------------------------


def _student_name(ws, row: int, name_col: int | None, first_col: int | None, last_col: int | None, code: str) -> str:
    if first_col and last_col:
        first = clean_text(ws.cell(row=row, column=first_col).value).upper()
        last = clean_text(ws.cell(row=row, column=last_col).value).upper()
        return f"{last}, {first}" if last else first
    if name_col:
        return clean_text(ws.cell(row=row, column=name_col).value).upper()
    return code


def parse_student_list(filepath: Path) -> tuple[str, list[dict[str, str]]]:
    """Read a Lista workbook (or CSV) into ``(section, students)``.

    Students are ``{"code", "name"}``; the section is inferred from the filename.
    """
    filepath = Path(filepath)
    section = filepath.stem.split("-")[0].strip()

    if filepath.suffix.lower() == ".csv":
        return section, _parse_student_csv(filepath)

    ws = load_workbook(filepath, data_only=True).worksheets[0]
    header_row = find_header_row_ws(ws)
    code_col = find_col_ws(ws, header_row, ["Codigo", "Code", "Student ID"])
    name_col = find_col_ws(ws, header_row, ["Nombre Completo", "Apellidos y Nombres", "Name"], required=False)
    first_col = find_col_ws(ws, header_row, ["First Name"], required=False)
    last_col = find_col_ws(ws, header_row, ["Last Name"], required=False)

    students = []
    for row in range(header_row + 1, ws.max_row + 1):
        code = normalize_student_code(ws.cell(row=row, column=code_col).value)
        if not code:
            continue
        students.append({"code": code, "name": _student_name(ws, row, name_col, first_col, last_col, code)})
    return section, students


def _parse_student_csv(filepath: Path) -> list[dict[str, str]]:
    rows = read_delimited(filepath)
    if not rows:
        return []

    columns = {normalize_header(column): column for column in rows[0].keys()}

    def pick(*needles: str) -> str | None:
        for needle in needles:
            for normalized, original in columns.items():
                if needle in normalized:
                    return original
        return None

    code_col = pick("codigo", "code", "student id")
    if not code_col:
        raise ValueError("Could not find a Code/Codigo column in the student list.")
    full_col = pick("nombre completo")
    first_col = pick("first name")
    last_col = pick("last name")
    fallback_col = pick("name", "nombre")

    students = []
    for row in rows:
        code = normalize_student_code(row.get(code_col))
        if not code:
            continue
        if first_col and last_col:
            first = clean_text(row.get(first_col)).upper()
            last = clean_text(row.get(last_col)).upper()
            name = f"{last}, {first}" if last else first
        elif full_col:
            name = clean_text(row.get(full_col)).upper()
        elif fallback_col:
            name = clean_text(row.get(fallback_col)).upper()
        else:
            name = code
        students.append({"code": code, "name": name})
    return students


def lastname_initial(student: dict) -> str:
    """First letter of the surname, for the "find a student" pickers."""
    name = clean_text(student.get("name", ""))
    lastname = name.split(",", 1)[0].strip() if "," in name else name
    return lastname[:1].upper()


def search_by_lastname_initial(
    students: list[dict],
    initial: str,
    excluded_codes: set[str] | None = None,
) -> list[dict]:
    excluded_codes = excluded_codes or set()
    initial = clean_text(initial).upper()[:1]
    if not initial:
        return []
    return [
        student
        for student in students
        if lastname_initial(student) == initial
        and normalize_student_code(student.get("code")) not in excluded_codes
    ]
