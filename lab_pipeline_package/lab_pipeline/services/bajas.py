"""Route 2 helper: moving a withdrawn student's row out of Lista/Notas.

Kept decoupled from the fixed-groups persistence in ``services/groups.py``:
this module never touches that JSON. A stale code left behind in a saved
fixed grouping is filtered out where it's read, not pruned here.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from .. import gradebook, naming, rosters
from ..common import append_row_dict, clean_text, normalize_student_code


def preview_withdrawal(course_root: Path, section: str, code: str) -> dict[str, object]:
    """Read-only: what withdrawing ``code`` would move, for a confirmation prompt."""
    course_root = Path(course_root)
    code = normalize_student_code(code)

    _, students = rosters.parse_student_list(naming.lista_path(course_root, section))
    student = next((item for item in students if item["code"] == code), None)
    if not student:
        raise ValueError(f"El codigo {code} no esta en la lista de la seccion {section}.")

    recorded_labs: list[int] = []
    notas_path = naming.notas_path(course_root, section)
    if notas_path.exists():
        book = gradebook.Gradebook(notas_path, read_only=True)
        recorded_labs = book.recorded_lab_numbers(code)

    return {"code": code, "name": student["name"], "recorded_labs": recorded_labs}


def _append_or_create(path: Path, row_dict: dict[str, object]) -> None:
    """Append ``row_dict`` to the Bajas workbook at ``path``, creating it (with
    ``row_dict``'s own keys as headers) if it doesn't exist yet."""
    path = Path(path)
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        for col, header in enumerate(row_dict.keys(), start=1):
            ws.cell(row=1, column=col).value = header

    append_row_dict(ws, header_row=1, row_dict=row_dict)
    wb.save(path)


def withdraw_student(course_root: Path, section: str, code: str) -> dict[str, object]:
    """Move ``code``'s row out of Lista (and Notas, if it exists) into Bajas."""
    course_root = Path(course_root)
    section = clean_text(section)
    code = normalize_student_code(code)

    lista_path = naming.lista_path(course_root, section)
    lista_row = rosters.pop_student_row(lista_path, code)
    if lista_row is None:
        raise ValueError(f"El codigo {code} no esta en la lista de la seccion {section}.")
    _append_or_create(naming.bajas_lista_path(course_root, section), lista_row)

    notas_row = None
    notas_path = naming.notas_path(course_root, section)
    if notas_path.exists():
        with gradebook.Gradebook(notas_path) as book:
            notas_row = book.pop_student(code)
        if notas_row is not None:
            _append_or_create(naming.bajas_notas_path(course_root, section), notas_row)

    return {
        "course_root": course_root,
        "section": section,
        "code": code,
        "name": lista_row.get("Nombre Completo", ""),
        "lista_bajas_path": naming.bajas_lista_path(course_root, section),
        "notas_bajas_path": naming.bajas_notas_path(course_root, section) if notas_row is not None else None,
        "success": True,
    }
