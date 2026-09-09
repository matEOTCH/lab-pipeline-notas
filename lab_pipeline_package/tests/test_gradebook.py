from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from lab_pipeline import gradebook
from lab_pipeline.common import find_header_row_ws


def header_names(ws, header_row: int) -> list[str]:
    return [ws.cell(row=header_row, column=col).value for col in range(1, ws.max_column + 1)]


def test_column_names_follow_the_lab_number():
    assert gradebook.group_column(2) == "Grupo_Lab2"
    assert gradebook.tm_percentage_column(2) == "%TM2"
    assert gradebook.responded_column(2) == "TM2_Resp?"
    assert gradebook.grade_column(2) == "Nota_Lab2"
    assert gradebook.block_columns(2, ["BB_Lab2"]) == ["BB_Lab2", "%TM2", "TM2_Resp?", "Nota_Lab2"]


def test_header_colors_are_distinct_per_role():
    assert gradebook.header_color("Grupo_Lab1", 1) == gradebook.GROUP_COLOR
    assert gradebook.header_color("Nota_Lab1", 1) == gradebook.GRADE_COLOR
    assert gradebook.header_color("TM1_Resp?", 1) == gradebook.RESPONDED_COLOR
    assert gradebook.header_color("BB_Lab1", 1) == gradebook.DEFAULT_COLOR


def test_normalize_custom_base_name_strips_spaces_and_trailing_digits():
    assert gradebook.normalize_custom_base_name("Quiz Lab 2") == "Quiz_Lab_"
    assert gradebook.normalize_custom_base_name("BB_Lab3") == "BB_Lab"
    assert gradebook.materialize_custom_column("BB_Lab3", 5) == "BB_Lab5"


def test_ensure_columns_creates_the_block_in_order(notas_path: Path):
    with gradebook.Gradebook(notas_path) as book:
        positions, created = book.ensure_columns(2, ["BB_Lab2"])

    assert created == ["Grupo_Lab2", "BB_Lab2", "%TM2", "TM2_Resp?", "Nota_Lab2"]

    ws = load_workbook(notas_path).worksheets[0]
    names = header_names(ws, find_header_row_ws(ws))
    assert names[-5:] == ["Grupo_Lab2", "BB_Lab2", "%TM2", "TM2_Resp?", "Nota_Lab2"]
    assert positions["Grupo_Lab2"] < positions["Nota_Lab2"]


def test_ensure_columns_is_idempotent(notas_path: Path):
    with gradebook.Gradebook(notas_path) as book:
        book.ensure_columns(2, ["BB_Lab2"])
    ws = load_workbook(notas_path).worksheets[0]
    first = header_names(ws, find_header_row_ws(ws))

    with gradebook.Gradebook(notas_path) as book:
        _, created = book.ensure_columns(2, ["BB_Lab2"])

    ws = load_workbook(notas_path).worksheets[0]
    assert created == []
    assert header_names(ws, find_header_row_ws(ws)) == first


def test_ensure_columns_reuses_an_existing_alias_instead_of_duplicating(notas_path: Path):
    wb = load_workbook(notas_path)
    ws = wb.worksheets[0]
    ws.cell(row=1, column=ws.max_column + 1).value = "Grupo Lab 2"
    wb.save(notas_path)

    with gradebook.Gradebook(notas_path) as book:
        _, created = book.ensure_columns(2)

    ws = load_workbook(notas_path).worksheets[0]
    names = header_names(ws, find_header_row_ws(ws))
    assert "Grupo_Lab2" not in created
    assert names.count("Grupo Lab 2") == 1
    assert "Grupo_Lab2" not in names


def test_two_labs_keep_separate_blocks(notas_path: Path):
    with gradebook.Gradebook(notas_path) as book:
        book.ensure_columns(1)
    with gradebook.Gradebook(notas_path) as book:
        book.ensure_columns(2)

    ws = load_workbook(notas_path).worksheets[0]
    names = header_names(ws, find_header_row_ws(ws))
    for lab in (1, 2):
        for name in ["Grupo_Lab", "%TM", "Nota_Lab"]:
            assert sum(1 for header in names if header and str(lab) in header and name.rstrip("_") in header) >= 1
    assert names.count("Nota_Lab1") == 1
    assert names.count("Nota_Lab2") == 1


def test_read_students_skips_blank_rows(notas_path: Path):
    book = gradebook.Gradebook(notas_path, read_only=True)

    assert len(book.students) == 6
    assert book.students[0]["code"] == "20250001"
    assert book.students[0]["name"] == "ALVAREZ, ANA"
    assert book.by_code["20250002"]["name"] == "BENITEZ, BRUNO"


def test_apply_grades_and_group_lookup_round_trip(notas_path: Path):
    with gradebook.Gradebook(notas_path) as book:
        positions, _ = book.ensure_columns(1)
        updated, unmatched = gradebook.write_group_numbers(
            book.ws, book.students, positions["Grupo_Lab1"], {"20250001": 3, "20250002": 3}
        )
        written = gradebook.apply_grades_by_code(
            book.ws, book.students, positions["Nota_Lab1"], {"20250001": 18, "99999999": 5}
        )

    assert (updated, written) == (2, 1)
    assert sorted(unmatched) == ["20250003", "20250004", "20250005", "20250006"]

    book = gradebook.Gradebook(notas_path, read_only=True)
    assert book.group_lookup(1) == {"20250001": 3, "20250002": 3}
    grade_col = gradebook.find_column(book.ws, book.header_row, "Nota_Lab1", 1)
    assert book.ws.cell(row=book.by_code["20250001"]["row"], column=grade_col).value == 18


def test_mark_absent_zeroes_the_block_and_formats_the_percentage(notas_path: Path):
    with gradebook.Gradebook(notas_path) as book:
        positions, _ = book.ensure_columns(2, ["BB_Lab2"])
        changed = gradebook.mark_absent(book.ws, book.by_code["20250003"], positions, 2, ["BB_Lab2"])

    assert [change["column"] for change in changed] == ["BB_Lab2", "Nota_Lab2", "%TM2", "TM2_Resp?"]

    book = gradebook.Gradebook(notas_path, read_only=True)
    row = book.by_code["20250003"]["row"]
    values = {
        name: book.ws.cell(row=row, column=gradebook.find_column(book.ws, book.header_row, name, 2)).value
        for name in ("BB_Lab2", "Nota_Lab2", "%TM2", "TM2_Resp?")
    }
    assert values == {"BB_Lab2": 0, "Nota_Lab2": 0, "%TM2": 0, "TM2_Resp?": "No"}

    wb = load_workbook(notas_path)
    ws = wb.worksheets[0]
    tm_col = gradebook.find_column(ws, book.header_row, "%TM2", 2)
    assert ws.cell(row=row, column=tm_col).number_format == gradebook.TM_NUMBER_FORMAT


def test_gradebook_context_manager_does_not_save_on_error(notas_path: Path):
    before = notas_path.read_bytes()
    try:
        with gradebook.Gradebook(notas_path) as book:
            book.ensure_columns(4)
            raise RuntimeError("boom")
    except RuntimeError:
        pass

    assert notas_path.read_bytes() == before
