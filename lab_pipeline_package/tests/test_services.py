from __future__ import annotations

from pathlib import Path

import pytest

from lab_pipeline import gradebook, naming, services, settings
from tests.test_teammates_csv import write_export


# --- structure ----------------------------------------------------------------


def test_create_structure_writes_lista_notas_and_remembers_the_domain(workspace: Path, roster_csv: Path):
    result = services.create_structure(
        workspace_root=workspace,
        project_name="2026",
        course_cycle_name="Quimica General 2026-1",
        section="315",
        email_domain="@aloe.ulima.edu.pe",
        csv_path=roster_csv,
    )

    assert result["list_path"].exists()
    assert result["notas_path"].exists()
    assert len(result["students"]) == 6
    assert result["warnings"] == []
    assert settings.project_settings(result["project_root"]).load()["email_domain"] == "aloe.ulima.edu.pe"
    assert naming.sections_in(result["course_root"]) == ["315"]


def test_create_structure_refuses_to_clobber_without_overwrite(workspace: Path, roster_csv: Path):
    kwargs = dict(
        workspace_root=workspace,
        project_name="2026",
        course_cycle_name="Quimica General 2026-1",
        section="315",
        email_domain="aloe.ulima.edu.pe",
        csv_path=roster_csv,
    )
    services.create_structure(**kwargs)

    with pytest.raises(FileExistsError):
        services.create_structure(**kwargs)

    result = services.create_structure(**kwargs, overwrite=True)
    assert result["warnings"] == ["Se reemplazaron archivos existentes de Lista/Notas."]


@pytest.mark.parametrize("field", ["project_name", "course_cycle_name", "section", "email_domain"])
def test_create_structure_rejects_blank_required_fields(workspace: Path, roster_csv: Path, field: str):
    kwargs = dict(
        workspace_root=workspace,
        project_name="2026",
        course_cycle_name="Quimica General 2026-1",
        section="315",
        email_domain="aloe.ulima.edu.pe",
        csv_path=roster_csv,
    )
    kwargs[field] = "   "

    with pytest.raises(ValueError):
        services.create_structure(**kwargs)


def test_create_structure_reports_a_missing_csv(workspace: Path):
    with pytest.raises(FileNotFoundError):
        services.create_structure(
            workspace_root=workspace,
            project_name="2026",
            course_cycle_name="Quimica General 2026-1",
            section="315",
            email_domain="aloe.ulima.edu.pe",
            csv_path=workspace / "nope.csv",
        )


# --- groups -------------------------------------------------------------------


@pytest.fixture
def section(workspace: Path, roster_csv: Path) -> dict:
    return services.create_structure(
        workspace_root=workspace,
        project_name="2026",
        course_cycle_name="Quimica General 2026-1",
        section="315",
        email_domain="aloe.ulima.edu.pe",
        csv_path=roster_csv,
    )


def test_create_lab_groups_writes_every_output_and_records_groups(section: dict):
    result = services.create_lab_groups(
        course_root=section["course_root"],
        section="315",
        professor="Perez, Martin",
        jefe="Tapia, Mateo",
        lab_number=2,
        group_count=3,
        seed=11,
    )

    for key in ("pdf_path", "groups_csv", "members_csv", "zip_path"):
        assert result[key].exists(), key
    assert sum(result["group_sizes"]) == 6
    assert result["notas_result"]["updated_count"] == 6
    assert result["notas_result"]["unmatched_codes"] == []
    assert naming.labs_in(section["course_root"], "315") == [2]

    book = gradebook.Gradebook(result["notas_path"], read_only=True)
    assert set(book.group_lookup(2).values()) == {1, 2, 3}


def test_create_lab_groups_remembers_professor_and_jefe(section: dict):
    services.create_lab_groups(
        course_root=section["course_root"],
        section="315",
        professor="Perez, Martin",
        jefe="Tapia, Mateo",
        lab_number=1,
        group_count=2,
        seed=1,
    )

    people = services.groups.section_people(section["course_root"], "315")
    assert people == {"professor": "Perez, Martin", "jefe": "Tapia, Mateo"}


def test_create_lab_groups_honours_forced_assignments(section: dict):
    result = services.create_lab_groups(
        course_root=section["course_root"],
        section="315",
        professor="P",
        jefe="J",
        lab_number=1,
        group_count=3,
        forced_assignments={"20250001": 2},
        seed=4,
    )

    book = gradebook.Gradebook(result["notas_path"], read_only=True)
    assert book.group_lookup(1)["20250001"] == 2


def test_regenerating_a_lab_overwrites_rather_than_duplicates(section: dict):
    kwargs = dict(
        course_root=section["course_root"],
        section="315",
        professor="P",
        jefe="J",
        lab_number=1,
        group_count=2,
    )
    services.create_lab_groups(**kwargs, seed=1)
    result = services.create_lab_groups(**kwargs, seed=2)

    book = gradebook.Gradebook(result["notas_path"], read_only=True)
    headers = [book.ws.cell(row=book.header_row, column=col).value for col in range(1, book.ws.max_column + 1)]
    assert headers.count("Grupo_Lab1") == 1


def test_create_lab_groups_needs_a_lista(section: dict):
    with pytest.raises(FileNotFoundError):
        services.create_lab_groups(
            course_root=section["course_root"],
            section="999",
            professor="P",
            jefe="J",
            lab_number=1,
        )


@pytest.mark.parametrize(("field", "value"), [("professor", ""), ("jefe", ""), ("section", "")])
def test_create_lab_groups_rejects_blank_people(section: dict, field: str, value: str):
    kwargs = dict(
        course_root=section["course_root"],
        section="315",
        professor="P",
        jefe="J",
        lab_number=1,
    )
    kwargs[field] = value

    with pytest.raises(ValueError):
        services.create_lab_groups(**kwargs)


def test_preview_groups_touches_no_files(section: dict):
    from lab_pipeline import rosters

    _, students = rosters.parse_student_list(naming.lista_path(section["course_root"], "315"))

    groups = services.preview_groups(students, 3, seed=9)

    assert sum(len(group) for group in groups) == 6
    assert naming.labs_in(section["course_root"], "315") == []


# --- grades -------------------------------------------------------------------


@pytest.fixture
def lab(section: dict) -> dict:
    return services.create_lab_groups(
        course_root=section["course_root"],
        section="315",
        professor="P",
        jefe="J",
        lab_number=1,
        group_count=3,
        seed=11,
    )


def column_value(notas_path: Path, code: str, column: str, lab_number: int):
    book = gradebook.Gradebook(notas_path, read_only=True)
    col = gradebook.find_column(book.ws, book.header_row, column, lab_number)
    return book.ws.cell(row=book.by_code[code]["row"], column=col).value


def test_configure_grade_columns_persists_and_materializes(section: dict):
    result = services.configure_grade_columns(
        section["course_root"], "315", 3, [{"base_name": "Quiz Lab", "method": "2"}]
    )

    assert result["materialized_columns"] == ["Quiz_Lab3"]
    assert services.grade_column_templates(section["course_root"], "315") == [
        {"base_name": "Quiz_Lab", "method": "2"}
    ]


def test_configure_grade_columns_rejects_a_bad_method(section: dict):
    with pytest.raises(ValueError, match="Metodo"):
        services.configure_grade_columns(
            section["course_root"], "315", 1, [{"base_name": "Quiz", "method": "9"}]
        )


def test_apply_grade_updates_writes_manual_student_grades(lab: dict):
    result = services.apply_grade_updates(
        notas_path=lab["notas_path"],
        lab_number=1,
        custom_grade_templates=[{"base_name": "BB_Lab", "method": "2"}],
        manual_student_grades={"BB_Lab1": {"20250001": "18", "20250002": "12,5"}},
    )

    assert result["custom_updates"] == {"BB_Lab1": 2}
    assert column_value(lab["notas_path"], "20250001", "BB_Lab1", 1) == 18
    assert column_value(lab["notas_path"], "20250002", "BB_Lab1", 1) == 12.5


def test_manual_group_grades_reach_every_member_of_the_group(lab: dict):
    book = gradebook.Gradebook(lab["notas_path"], read_only=True)
    lookup = book.group_lookup(1)
    member = next(code for code, group in lookup.items() if group == 1)

    services.apply_grade_updates(
        notas_path=lab["notas_path"],
        lab_number=1,
        custom_grade_templates=[{"base_name": "Inf_Lab", "method": "3"}],
        manual_group_grades={"Inf_Lab1": {"1": 16}},
    )

    assert column_value(lab["notas_path"], member, "Inf_Lab1", 1) == 16
    other = next(code for code, group in lookup.items() if group == 2)
    assert column_value(lab["notas_path"], other, "Inf_Lab1", 1) is None


def test_payload_may_be_keyed_by_index_base_name_or_column_name(lab: dict):
    """The browser must not have to reproduce the server's naming rules."""
    for key in ("0", "Quiz Lab", "Quiz_Lab1"):
        services.apply_grade_updates(
            notas_path=lab["notas_path"],
            lab_number=1,
            custom_grade_templates=[{"base_name": "Quiz Lab", "method": "2"}],
            manual_student_grades={key: {"20250001": 15}},
        )
        assert column_value(lab["notas_path"], "20250001", "Quiz_Lab1", 1) == 15


def test_blackboard_file_column_uses_the_uploaded_file(lab: dict, tmp_path: Path):
    grade_file = tmp_path / "bb.csv"
    grade_file.write_text(
        "Apellidos,Nombre,Nombre de usuario,Quiz\nAlvarez,Ana,20250001,17\n", encoding="utf-8"
    )

    services.apply_grade_updates(
        notas_path=lab["notas_path"],
        lab_number=1,
        custom_grade_templates=[{"base_name": "BB_Lab", "method": "1"}],
        uploaded_files={"BB_Lab1": grade_file},
    )

    assert column_value(lab["notas_path"], "20250001", "BB_Lab1", 1) == 17


def test_a_missing_blackboard_file_is_reported(lab: dict):
    with pytest.raises(ValueError, match="Falta archivo"):
        services.apply_grade_updates(
            notas_path=lab["notas_path"],
            lab_number=1,
            custom_grade_templates=[{"base_name": "BB_Lab", "method": "1"}],
        )


def test_teammates_import_fills_percentage_response_and_comments(lab: dict, tmp_path: Path):
    export_path = write_export(tmp_path / "tm.csv")

    result = services.apply_grade_updates(
        notas_path=lab["notas_path"],
        lab_number=1,
        teammates_csv_path=export_path,
    )

    assert result["tm_result"]["updated_tm"] == 2
    assert column_value(lab["notas_path"], "20250001", "%TM1", 1) == 1.0
    assert column_value(lab["notas_path"], "20250002", "%TM1", 1) == 0.85
    assert column_value(lab["notas_path"], "20250001", "TM1_Resp?", 1) == "Si"
    assert column_value(lab["notas_path"], "20250003", "TM1_Resp?", 1) == "No"

    from openpyxl import load_workbook

    wb = load_workbook(lab["notas_path"])
    assert "TM_Comments per student" in wb.sheetnames
    comments = wb["TM_Comments per student"]
    headers = [comments.cell(row=1, column=col).value for col in range(1, comments.max_column + 1)]
    assert "Q2_Lab1" in headers and "Q3_Lab1" in headers


def test_absences_zero_the_whole_block_in_the_same_pass(lab: dict):
    result = services.apply_grade_updates(
        notas_path=lab["notas_path"],
        lab_number=1,
        custom_grade_templates=[{"base_name": "BB_Lab", "method": "2"}],
        manual_student_grades={"BB_Lab1": {"20250003": 20}},
        absent_codes=["20250003"],
    )

    assert result["absences"]["absences_updated"] == 1
    for column, expected in (("BB_Lab1", 0), ("Nota_Lab1", 0), ("%TM1", 0), ("TM1_Resp?", "No")):
        assert column_value(lab["notas_path"], "20250003", column, 1) == expected


def test_apply_absences_alone_reports_unknown_codes(lab: dict):
    result = services.apply_absences(lab["notas_path"], 1, ["20250004", "99999999"])

    assert result["absences_updated"] == 1
    assert result["not_found_codes"] == ["99999999"]


def test_apply_absences_with_no_codes_is_a_no_op(lab: dict):
    before = lab["notas_path"].read_bytes()

    result = services.apply_absences(lab["notas_path"], 1, [])

    assert result["absences_updated"] == 0
    assert lab["notas_path"].read_bytes() == before


def test_preview_absences_changes_nothing(lab: dict):
    before = lab["notas_path"].read_bytes()

    result = services.preview_absences(lab["notas_path"], 1, ["20250001", "99999999"], ["BB_Lab1"])

    assert result["not_found_codes"] == ["99999999"]
    assert [student["code"] for student in result["students"]] == ["20250001"]
    assert result["target_columns"] == ["BB_Lab1", "Nota_Lab1", "%TM1", "TM1_Resp?"]
    assert lab["notas_path"].read_bytes() == before


def test_a_failure_mid_request_leaves_the_workbook_untouched(lab: dict):
    """The old code saved between stages, so a late error left it half-written."""
    before = lab["notas_path"].read_bytes()

    with pytest.raises(ValueError):
        services.apply_grade_updates(
            notas_path=lab["notas_path"],
            lab_number=1,
            custom_grade_templates=[
                {"base_name": "Ok_Lab", "method": "2"},
                {"base_name": "Broken_Lab", "method": "1"},
            ],
            manual_student_grades={"Ok_Lab1": {"20250001": 15}},
        )

    assert lab["notas_path"].read_bytes() == before


# --- bajas ----------------------------------------------------------------


def test_withdraw_student_moves_lista_and_notas_rows(section: dict):
    """``section`` already has a Notas workbook - create_structure seeds it from Lista."""
    from lab_pipeline import rosters

    result = services.withdraw_student(section["course_root"], "315", "20250003")

    assert result["success"] is True
    assert result["name"] == "CASTRO, CARLA"
    assert result["lista_bajas_path"].exists()
    assert result["notas_bajas_path"] is not None
    assert result["notas_bajas_path"].exists()

    _, remaining = rosters.parse_student_list(naming.lista_path(section["course_root"], "315"))
    assert "20250003" not in [student["code"] for student in remaining]

    _, bajas_students = rosters.parse_student_list(result["lista_bajas_path"])
    assert [student["code"] for student in bajas_students] == ["20250003"]

    book = gradebook.Gradebook(naming.notas_path(section["course_root"], "315"), read_only=True)
    assert "20250003" not in book.by_code


def test_withdraw_student_with_no_notas_workbook_yet(workspace: Path, roster_csv: Path):
    """When Notas doesn't exist at all, only the Lista side moves."""
    from lab_pipeline import rosters

    result = services.create_structure(
        workspace_root=workspace,
        project_name="2026",
        course_cycle_name="Quimica General 2026-1",
        section="315",
        email_domain="aloe.ulima.edu.pe",
        csv_path=roster_csv,
    )
    result["notas_path"].unlink()

    withdrawal = services.withdraw_student(result["course_root"], "315", "20250003")

    assert withdrawal["lista_bajas_path"].exists()
    assert withdrawal["notas_bajas_path"] is None

    _, remaining = rosters.parse_student_list(naming.lista_path(result["course_root"], "315"))
    assert "20250003" not in [student["code"] for student in remaining]


def test_withdraw_student_moves_notas_row_too_when_it_exists(lab: dict, section: dict):
    with gradebook.Gradebook(lab["notas_path"]) as book:
        positions, _ = book.ensure_columns(1)
        gradebook.apply_grades_by_code(book.ws, book.students, positions["Nota_Lab1"], {"20250002": 17})

    result = services.withdraw_student(section["course_root"], "315", "20250002")

    assert result["notas_bajas_path"] is not None
    assert result["notas_bajas_path"].exists()

    book = gradebook.Gradebook(lab["notas_path"], read_only=True)
    assert "20250002" not in book.by_code

    from openpyxl import load_workbook

    bajas_ws = load_workbook(result["notas_bajas_path"]).active
    headers = [bajas_ws.cell(row=1, column=col).value for col in range(1, bajas_ws.max_column + 1)]
    row = dict(zip(headers, [bajas_ws.cell(row=2, column=col).value for col in range(1, bajas_ws.max_column + 1)]))
    assert row["Codigo"] == "20250002"
    assert row["Nota_Lab1"] == 17


def test_withdraw_unknown_code_raises(section: dict):
    with pytest.raises(ValueError):
        services.withdraw_student(section["course_root"], "315", "99999999")


def test_preview_withdrawal_reports_recorded_labs(lab: dict, section: dict):
    with gradebook.Gradebook(lab["notas_path"]) as book:
        positions, _ = book.ensure_columns(1)
        gradebook.apply_grades_by_code(book.ws, book.students, positions["Nota_Lab1"], {"20250001": 18})

    preview = services.preview_withdrawal(section["course_root"], "315", "20250001")
    assert preview == {"code": "20250001", "name": "ALVAREZ, ANA", "recorded_labs": [1]}

    other_preview = services.preview_withdrawal(section["course_root"], "315", "20250002")
    assert other_preview["recorded_labs"] == []


def test_preview_withdrawal_rejects_an_unknown_code(section: dict):
    with pytest.raises(ValueError):
        services.preview_withdrawal(section["course_root"], "315", "99999999")


def test_withdrawal_reconciles_notas_headers_across_two_labs(section: dict):
    """A second withdrawal, after more lab columns exist, must not drop the first row's columns."""
    from openpyxl import load_workbook

    services.create_lab_groups(
        course_root=section["course_root"], section="315", professor="P", jefe="J",
        lab_number=1, group_count=2, seed=1,
    )
    services.withdraw_student(section["course_root"], "315", "20250001")

    services.create_lab_groups(
        course_root=section["course_root"], section="315", professor="P", jefe="J",
        lab_number=2, group_count=2, seed=2,
    )
    services.withdraw_student(section["course_root"], "315", "20250002")

    ws = load_workbook(naming.bajas_notas_path(section["course_root"], "315")).active
    headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
    assert "Grupo_Lab1" in headers
    assert "Grupo_Lab2" in headers
    assert ws.max_row == 3


# --- fixed groups (groups_csv) ---------------------------------------------


def test_fixed_groups_round_trips_through_settings(section: dict):
    services.groups.remember_fixed_groups(section["course_root"], "315", {"20250001": 2, "20250002": 1})

    assert services.groups.fixed_groups(section["course_root"], "315") == {"20250001": 2, "20250002": 1}
    assert services.groups.fixed_groups(section["course_root"], "999") == {}


def test_create_lab_groups_places_students_as_the_imported_csv_specifies(section: dict, tmp_path: Path):
    from lab_pipeline import groups_csv

    csv_path = tmp_path / "grupos.csv"
    csv_path.write_text(
        "Group Code*,User Name*\n"
        "Grupo_gc_Lab1_1,20250001\n"
        "Grupo_gc_Lab1_1,20250002\n"
        "Grupo_gc_Lab1_2,20250003\n",
        encoding="utf-8",
    )
    forced = groups_csv.parse_groups_csv(csv_path)
    assert forced == {"20250001": 1, "20250002": 1, "20250003": 2}

    result = services.create_lab_groups(
        course_root=section["course_root"],
        section="315",
        professor="P",
        jefe="J",
        lab_number=1,
        group_count=2,
        forced_assignments=forced,
        seed=3,
    )

    book = gradebook.Gradebook(result["notas_path"], read_only=True)
    lookup = book.group_lookup(1)
    assert lookup["20250001"] == 1
    assert lookup["20250002"] == 1
    assert lookup["20250003"] == 2


# --- status -------------------------------------------------------------------


def test_course_status_reports_sections_labs_and_outputs(lab: dict, section: dict):
    status = services.course_status(section["course_root"])

    assert status["sections"] == ["315"]
    section_row = status["section_status"][0]
    assert section_row["labs"] == [1]
    assert section_row["notas_path"] is not None
    assert section_row["outputs"][1]["zip"] is not None
